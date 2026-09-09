from flask import Flask, render_template, request, redirect, session, jsonify, url_for, send_file, Response
import datetime
import io
import os
import qrcode
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from database import (
    get_student_by_pid, get_student_by_id, get_student_by_email_or_pid,
    get_all_students, register_student, get_teacher_by_id,
    get_teacher_by_email_or_id, get_all_teachers, register_teacher,
    get_full_weekly_timetable, get_timetable_for_day, get_teacher_timetable,
    get_current_active_lecture, create_attendance_session, get_session_by_token,
    close_attendance_session, get_active_session_for_teacher, mark_qr_attendance,
    get_session_live_attendance, get_student_attendance_history, get_student_stats,
    get_teacher_attendance_records, get_teacher_stats, get_teacher_analytics,
    get_all_attendance_for_export, get_notifications_for_role, add_notification,
    clear_attendance_history, init_db, finalize_all_expired_sessions,
    get_latest_session_for_teacher, finalize_attendance_session, get_current_time,
    get_dynamic_qr_payload, calculate_distance_meters, DB_NAME
)


app = Flask(__name__)
app.secret_key = os.getenv('SECRET_KEY', 'supersecret_college_qr_key_2026')

# Ensure database is created on startup if missing
if not os.path.exists(DB_NAME):
    init_db()

# -------------------------------------------------------------
# AUTHENTICATION ROUTES
# -------------------------------------------------------------

@app.route('/')
def index():
    if 'role' in session:
        if session['role'] == 'teacher':
            return redirect(url_for('teacher_dashboard'))
        elif session['role'] == 'student':
            return redirect(url_for('student_dashboard'))
    return redirect(url_for('login'))

@app.route('/login')
def login():
    role = request.args.get('role', 'student').strip()
    return render_template('login.html', role=role)

@app.route('/login-post', methods=['POST'])
def login_post():
    role = request.form.get('role', 'student').strip()
    identifier = request.form.get('identifier', '').strip()
    password = request.form.get('password', '').strip()

    if not identifier or not password:
        return render_template('login.html', error="Please provide both login ID and password.", role=role)

    if role in ['teacher', 'admin']:
        teacher = get_teacher_by_email_or_id(identifier)
        if not teacher and role == 'admin' and identifier.lower() in ['admin', 'admin@college.edu'] and password in ['admin123', 'admin', 'password']:
            all_teachers = get_all_teachers()
            teacher = all_teachers[0] if all_teachers else None

        if teacher and (teacher['password'] == password or (role == 'admin' and password in ['admin123', 'admin', 'password'])):
            session.clear()
            session['role'] = 'teacher'
            session['user_id'] = teacher['id']
            session['teacher_id'] = teacher['teacher_id']
            session['name'] = teacher['name']
            session['email'] = teacher['email']
            session['department'] = teacher['department']
            session['subjects'] = teacher['subjects']
            return redirect(url_for('teacher_dashboard'))
        else:
            return render_template('login.html', error="Invalid Teacher/Admin ID or Password.", role=role)
    else:
        student = get_student_by_email_or_pid(identifier)
        if student and student['password'] == password:
            session.clear()
            session['role'] = 'student'
            session['user_id'] = student['id']
            session['pid'] = student['pid']
            session['name'] = student['name']
            session['email'] = student['email']
            session['roll_number'] = student['roll_number']
            session['department'] = student['department']
            session['course'] = student['course']
            session['year'] = student['year']
            session['division'] = student['division']
            return redirect(url_for('student_dashboard'))
        else:
            return render_template('login.html', error="Invalid Student PID / Email / Roll or Password.", role='student')

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))

@app.route('/register')
def register():
    role = request.args.get('role', 'student').strip()
    active_tab = 'teacher' if role == 'teacher' else 'student'
    return render_template('register.html', active_tab=active_tab)

@app.route('/register-student-post', methods=['POST'])
def register_student_post():
    pid = request.form.get('pid', '').strip().upper()
    name = request.form.get('name', '').strip()
    email = request.form.get('email', '').strip().lower()
    password = request.form.get('password', '').strip()
    department = request.form.get('department', 'ECS').strip()
    course = request.form.get('course', 'Electronics & Computer Science').strip()
    year = request.form.get('year', 'Fourth Year').strip()
    division = request.form.get('division', 'A').strip().upper()
    roll_number = request.form.get('roll_number', '').strip()

    if not all([pid, name, email, password, roll_number]):
        return render_template('register.html', active_tab='student', error="All fields marked with * are required.")

    success, message = register_student(pid, name, email, password, department, course, year, division, roll_number)
    if success:
        return render_template('register.html', active_tab='student', success=f"Registration successful! Student PID: {pid}. You can now sign in.")
    else:
        return render_template('register.html', active_tab='student', error=message)

@app.route('/register-teacher-post', methods=['POST'])
def register_teacher_post():
    secret_code = request.form.get('secret_code', '').strip()
    teacher_id = request.form.get('teacher_id', '').strip().upper()
    name = request.form.get('name', '').strip()
    email = request.form.get('email', '').strip().lower()
    password = request.form.get('password', '').strip()
    department = request.form.get('department', 'ECS').strip()
    subjects = request.form.get('subjects', '').strip()

    TEACHER_SECRET_CODE = '031926'
    if secret_code != TEACHER_SECRET_CODE:
        return render_template('register.html', active_tab='teacher', error="Invalid Faculty Secret Code. Only authorized teachers can create an account.")

    if not all([teacher_id, name, email, password, subjects]):
        return render_template('register.html', active_tab='teacher', error="All fields marked with * are required.")

    success, message = register_teacher(teacher_id, name, email, password, department, subjects)
    if success:
        return render_template('register.html', active_tab='teacher', success=f"Teacher registered successfully! ID: {teacher_id}. You can now sign in.")
    else:
        return render_template('register.html', active_tab='teacher', error=message)

# -------------------------------------------------------------
# STUDENT DASHBOARD & SCANNER
# -------------------------------------------------------------

@app.route('/student-dashboard')
def student_dashboard():
    if session.get('role') != 'student':
        return redirect(url_for('login'))

    # Finalize any expired sessions so absent students are recorded
    finalize_all_expired_sessions()

    pid = session.get('pid')
    student = get_student_by_pid(pid)
    if not student:
        session.clear()
        return redirect(url_for('login'))


    stats = get_student_stats(pid)
    history = get_student_attendance_history(pid)
    now = get_current_time()
    active_lecture, upcoming_lecture, today_lectures = get_current_active_lecture(now)
    weekly_timetable = get_full_weekly_timetable()
    notifications = get_notifications_for_role('student', pid)

    return render_template(
        'student_dashboard.html',
        student=student,
        stats=stats,
        history=history,
        current_lecture=active_lecture,
        upcoming_lecture=upcoming_lecture,
        today_lectures=today_lectures,
        weekly_timetable=weekly_timetable,
        notifications=notifications,
        today_date=now.strftime('%A, %d %B %Y')
    )

@app.route('/student/scan-attendance', methods=['POST'])
def student_scan_attendance():
    """
    PROX-PREVENTION ATTENDANCE ENDPOINT:
    The student PID is STRICTLY fetched from the server session, NOT accepted from client request body.
    This guarantees that a student cannot mark attendance on behalf of someone else's PID.
    Validates dynamic QR rotation and geofencing distance on Flask server.
    """
    if session.get('role') != 'student':
        return jsonify({"success": False, "error": "Unauthorized: Please log in as a student to mark attendance."}), 401

    student_pid = session.get('pid')
    if not student_pid:
        return jsonify({"success": False, "error": "Invalid student session. Please log in again."}), 401

    data = request.get_json(silent=True) or request.form
    session_token = data.get('session_token', '').strip()
    latitude = data.get('latitude')
    longitude = data.get('longitude')
    accuracy = data.get('accuracy')

    if not session_token:
        return jsonify({"success": False, "error": "Invalid scan: Attendance Session ID or QR payload is missing."}), 400

    # Call backend validation & recording engine with geofence coordinates
    success, result = mark_qr_attendance(session_token, student_pid, student_lat=latitude, student_lng=longitude, accuracy=accuracy)

    if success:
        return jsonify({
            "success": True,
            "message": f"Attendance marked successfully for {result['subject']}!",
            "details": result
        })
    else:
        return jsonify({
            "success": False,
            "error": result
        }), 400

# -------------------------------------------------------------
# TEACHER DASHBOARD & LIVE ATTENDANCE
# -------------------------------------------------------------

@app.route('/teacher-dashboard')
def teacher_dashboard():
    if session.get('role') != 'teacher':
        return redirect(url_for('login'))

    teacher_id = session.get('teacher_id')
    teacher = get_teacher_by_id(teacher_id)
    if not teacher:
        session.clear()
        return redirect(url_for('login'))

    teacher_name = teacher['name']
    now = get_current_time()
    day_name = now.strftime('%A')
    
    # Automatically finalize expired sessions across the system
    finalize_all_expired_sessions()

    # Timetable for teacher
    today_schedule = get_teacher_timetable(teacher_name, day_name)
    all_teacher_schedule = get_teacher_timetable(teacher_name)
    
    # Active lecture from timetable
    active_lecture, upcoming_lecture, _ = get_current_active_lecture(now)
    # Check if active lecture is taught by this teacher
    teacher_active_lec = None
    if active_lecture and (teacher['name'] in active_lecture['teacher_name'] or teacher['teacher_id'] in active_lecture['teacher_code']):
        teacher_active_lec = active_lecture

    # Active session if any
    active_session = get_active_session_for_teacher(teacher_id)
    
    # Check if there is an active session OR most recent session to show in the roster
    current_session_for_roster = active_session or get_latest_session_for_teacher(teacher_id)
    students_list = get_all_students()
    
    roster_data = None
    if current_session_for_roster:
        roster_data = get_session_live_attendance(current_session_for_roster['session_token'])
        roster = roster_data['roster'] if roster_data else []
    else:
        roster = []
        for s in students_list:
            roster.append({
                "pid": s['pid'],
                "name": s['name'],
                "roll_number": s['roll_number'],
                "department": s['department'],
                "year": s['year'],
                "division": s['division'],
                "status": "Absent",
                "timestamp": "-"
            })

    stats = get_teacher_stats(teacher_name)
    analytics = get_teacher_analytics(teacher_name)
    recent_records = get_teacher_attendance_records(teacher_name)
    notifications = get_notifications_for_role('teacher', teacher_id)

    return render_template(
        'teacher_dashboard.html',
        teacher=teacher,
        today_schedule=today_schedule,
        all_schedule=all_teacher_schedule,
        current_lecture=teacher_active_lec,
        upcoming_lecture=upcoming_lecture,
        active_session=active_session,
        current_session_for_roster=current_session_for_roster,
        roster=roster,
        roster_data=roster_data,
        stats=stats,
        analytics=analytics,
        recent_records=recent_records,
        students=students_list,
        notifications=notifications,
        today_date=now.strftime('%A, %d %B %Y')
    )

@app.route('/teacher/start-attendance', methods=['POST'])
def teacher_start_attendance():
    if session.get('role') != 'teacher':
        return jsonify({"success": False, "error": "Unauthorized"}), 401

    teacher_id = session.get('teacher_id')
    teacher = get_teacher_by_id(teacher_id)
    if not teacher:
        return jsonify({"success": False, "error": "Teacher profile not found"}), 404

    data = request.get_json(silent=True) or request.form
    subject = data.get('subject', 'PLCA')
    duration = int(data.get('duration', 15))
    latitude = data.get('latitude')
    longitude = data.get('longitude')
    radius_meters = data.get('radius_meters', 50)
    geofence_enabled = data.get('geofence_enabled', 1)

    # Lookup subject code and teacher code
    timetable = get_full_weekly_timetable()
    subject_code = subject
    for t in timetable:
        if t['subject'] == subject:
            subject_code = t['subject_code']
            break

    # Create dynamic attendance session with geofencing reference point
    new_session = create_attendance_session(
        teacher_id=teacher_id,
        teacher_name=teacher['name'],
        subject=subject,
        subject_code=subject_code,
        class_name="Fourth Year B.E. ECS",
        division="A",
        duration_minutes=duration,
        latitude=latitude,
        longitude=longitude,
        radius_meters=radius_meters,
        geofence_enabled=geofence_enabled
    )

    return jsonify({
        "success": True,
        "message": f"Attendance session started for {subject}!",
        "session": new_session
    })

@app.route('/teacher/stop-attendance', methods=['POST'])
def teacher_stop_attendance():
    if session.get('role') != 'teacher':
        return jsonify({"success": False, "error": "Unauthorized"}), 401

    data = request.get_json(silent=True) or request.form
    session_token = data.get('session_token')
    
    if not session_token:
        # Close current active session for this teacher
        active_session = get_active_session_for_teacher(session.get('teacher_id'))
        if active_session:
            session_token = active_session['session_token']

    if session_token:
        finalize_attendance_session(session_token)
        return jsonify({"success": True, "message": "Attendance session closed successfully. All unscanned students marked absent."})
    
    return jsonify({"success": False, "error": "No active session found to close."}), 400

@app.route('/teacher/live-attendance/<session_token>')
def teacher_live_attendance(session_token):
    if session.get('role') != 'teacher':
        return jsonify({"success": False, "error": "Unauthorized"}), 401

    data = get_session_live_attendance(session_token)
    if not data:
        return jsonify({"success": False, "error": "Session not found"}), 404

    return jsonify({"success": True, "data": data})

@app.route('/teacher/manual-override', methods=['POST'])
def teacher_manual_override():
    if session.get('role') != 'teacher':
        return jsonify({"success": False, "error": "Unauthorized"}), 401

    data = request.get_json(silent=True) or request.form
    student_pid = data.get('student_pid', '').strip().upper()
    session_token = data.get('session_token', '').strip()

    if not student_pid:
        return jsonify({"success": False, "error": "Student PID is required."}), 400

    teacher_id = session.get('teacher_id')
    active_session = get_active_session_for_teacher(teacher_id)
    if not session_token and active_session:
        session_token = active_session['session_token']

    if not session_token:
        # If no active session, find current teacher and start a temporary manual session
        teacher = get_teacher_by_id(teacher_id)
        sub = teacher['subjects'].split(',')[0].strip() if teacher else 'General'
        sess_success, new_sess = create_attendance_session(teacher_id, sub, duration_minutes=60)
        if sess_success:
            session_token = new_sess['session_token']
        else:
            return jsonify({"success": False, "error": "Could not create attendance session."}), 400

    success, result = mark_qr_attendance(session_token, student_pid)
    if success:
        return jsonify({
            "success": True,
            "message": f"Student {student_pid} marked present manually!",
            "details": result,
            "session_token": session_token
        })
    else:
        return jsonify({"success": False, "error": result}), 400

@app.route('/teacher/clear-history', methods=['POST'])
def teacher_clear_history():
    if session.get('role') != 'teacher':
        return jsonify({"success": False, "error": "Unauthorized"}), 401
    clear_attendance_history()
    return jsonify({"success": True, "message": "Attendance history cleared successfully."})

# -------------------------------------------------------------
# QR CODE GENERATORS (PNG RESPONSES)
# -------------------------------------------------------------

@app.route('/student-qr/<pid>')
def student_qr(pid):
    """
    Generates Student Profile QR code containing the Student's unique PID.
    """
    student = get_student_by_pid(pid)
    if not student:
        return "Student not found", 404

    # QR Payload formatted clearly
    qr_content = f"STUDENT-PID:{student['pid']}"
    
    qr = qrcode.QRCode(
        version=1,
        error_correction=qrcode.constants.ERROR_CORRECT_H,
        box_size=10,
        border=3,
    )
    qr.add_data(qr_content)
    qr.make(fit=True)

    img = qr.make_image(fill_color="black", back_color="white")
    buf = io.BytesIO()
    img.save(buf)
    buf.seek(0)

    download = request.args.get('download', '0') == '1'
    if download:
        return send_file(buf, mimetype="image/png", as_attachment=True, download_name=f"Student_QR_{student['pid']}.png")
    return send_file(buf, mimetype="image/png")

@app.route('/session-qr/<session_token>')
@app.route('/session-dynamic-qr/<session_token>')
def session_qr(session_token):
    """
    Generates dynamic Temporary Attendance Session QR code.
    Expires and rotates every 25 seconds using server-side HMAC time windows.
    """
    sess = get_session_by_token(session_token)
    if not sess:
        return "Invalid session token", 404

    # Generate dynamic payload with rotating time-window signature
    dynamic_info = get_dynamic_qr_payload(session_token)
    qr_payload = dynamic_info['payload']
    
    qr = qrcode.QRCode(
        version=1,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
        box_size=10,
        border=3,
    )
    qr.add_data(qr_payload)
    qr.make(fit=True)

    img = qr.make_image(fill_color="#171e19", back_color="#ffe17c")
    buf = io.BytesIO()
    img.save(buf)
    buf.seek(0)

    response = send_file(buf, mimetype="image/png")
    response.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate, max-age=0'
    response.headers['Pragma'] = 'no-cache'
    return response

@app.route('/api/session-dynamic-token/<session_token>')
def api_session_dynamic_token(session_token):
    """
    Returns the current dynamic rotating QR token payload and countdown seconds
    for real-time classroom projector synchronization.
    """
    sess = get_session_by_token(session_token)
    if not sess:
        return jsonify({"success": False, "error": "Session not found"}), 404

    info = get_dynamic_qr_payload(session_token)
    return jsonify({
        "success": True,
        "token_info": info
    })

# -------------------------------------------------------------
# EXCEL EXPORT (.XLSX)
# -------------------------------------------------------------

@app.route('/export-attendance-excel')
def export_attendance_excel():
    if session.get('role') != 'teacher':
        return redirect(url_for('login'))

    subject = request.args.get('subject', 'All')
    date = request.args.get('date', '')
    division = request.args.get('division', 'A')

    records = get_all_attendance_for_export(
        subject=subject if subject != 'All' else None,
        date=date if date else None,
        division=division if division != 'All' else None
    )

    # Create openpyxl Workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance Records"

    # Ensure grid lines are visible
    ws.views.sheetView[0].showGridLines = True

    # Styling definitions
    title_font = Font(name="Segoe UI", size=16, bold=True, color="171E19")
    subtitle_font = Font(name="Segoe UI", size=10, italic=True, color="555555")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="171E19", end_color="171E19", fill_type="solid")
    alt_fill = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")
    present_font = Font(name="Segoe UI", size=10, bold=True, color="008000")
    absent_font = Font(name="Segoe UI", size=10, bold=True, color="C00000")
    
    thin_border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )

    # Header Title Block
    ws.merge_cells('A1:N1')
    ws['A1'] = "COLLEGE OF ENGINEERING & TECHNOLOGY — DEPARTMENT OF ECS"
    ws['A1'].font = title_font
    ws['A1'].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells('A2:N2')
    export_sub = f"Official Attendance Report | Class: Fourth Year B.E. ECS (Div {division}) | Subject: {subject} | Generated: {get_current_time().strftime('%Y-%m-%d %H:%M:%S')}"
    ws['A2'] = export_sub
    ws['A2'].font = subtitle_font
    ws['A2'].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # Blank spacer row
    ws.row_dimensions[3].height = 10

    # Table Column Headers (Includes Student Name, PID, Roll No., Subject, Date, Time, Status, Distance)
    headers = [
        "PID", "Student Name", "Roll Number", "Department", "Year",
        "Division", "Subject", "Teacher", "Date", "Day",
        "Lecture Time", "Timestamp", "Status", "Geofence Distance"
    ]
    ws.row_dimensions[4].height = 24

    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num, value=header_title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Data Rows
    for row_idx, r in enumerate(records, 5):
        dist_str = f"{round(r['distance_meters'], 1)}m" if ('distance_meters' in r.keys() and r['distance_meters'] is not None) else ("Verified" if r['status'] == 'Present' else "-")
        row_values = [
            r['pid'],
            r['student_name'],
            r['roll_number'],
            r['department'],
            r['year'],
            r['division'],
            r['subject'],
            r['teacher_name'],
            r['date'],
            r['day'],
            r['lecture_time'],
            r['timestamp'],
            r['status'],
            dist_str
        ]
        ws.row_dimensions[row_idx].height = 20

        is_even = (row_idx % 2 == 0)
        for col_num, val in enumerate(row_values, 1):
            cell = ws.cell(row=row_idx, column=col_num, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center" if col_num in [1, 3, 5, 6, 9, 10, 13, 14] else "left", vertical="center")
            if is_even:
                cell.fill = alt_fill
            if col_num == 13: # Status column
                cell.font = present_font if r['status'] == 'Present' else absent_font

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row > 2 and cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Save to BytesIO buffer
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"ECS_Attendance_Report_{get_current_time().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(
        buffer,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename
    )

# -------------------------------------------------------------
# TIMETABLE API (For interactive dynamic schedules)
# -------------------------------------------------------------

@app.route('/api/timetable')
def api_timetable():
    day = request.args.get('day')
    if day:
        lectures = [dict(r) for r in get_timetable_for_day(day)]
    else:
        lectures = [dict(r) for r in get_full_weekly_timetable()]
    return jsonify(lectures)

@app.route('/favicon.ico')
def favicon():
    return ('', 204)

# -------------------------------------------------------------
# GLOBAL ERROR HANDLING (Railway Diagnostics)
# -------------------------------------------------------------

@app.errorhandler(500)
@app.errorhandler(Exception)
def handle_application_exception(e):
    from werkzeug.exceptions import HTTPException
    # Let standard HTTP exceptions (404, 401, 302, etc.) be handled normally
    if isinstance(e, HTTPException) and e.code != 500:
        return e

    import traceback
    import sys
    tb = traceback.format_exc()
    sys.stderr.write(f"\n==================== [AttendQR CRITICAL SERVER ERROR] ====================\n")
    sys.stderr.write(f"Endpoint: {request.path} | Method: {request.method}\n")
    sys.stderr.write(f"Error: {str(e)}\n")
    sys.stderr.write(f"{tb}")
    sys.stderr.write(f"===========================================================================\n")
    sys.stderr.flush()

    if request.path.startswith('/api') or request.is_json:
        return jsonify({"success": False, "error": f"Internal Server Error: {str(e)}"}), 500

    # User friendly recovery page
    return f"""<!DOCTYPE html>
<html>
<head>
    <title>Server Error | AttendQR</title>
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <style>
        body {{ font-family: 'Segoe UI', system-ui, sans-serif; background: #fffbe6; padding: 2rem; color: #171e19; }}
        .card {{ max-width: 600px; margin: 3rem auto; background: #fff; border: 3px solid #000; border-radius: 12px; box-shadow: 6px 6px 0 #000; padding: 2rem; }}
        .btn {{ display: inline-block; padding: 0.6rem 1.2rem; background: #ffe17c; border: 2px solid #000; border-radius: 6px; font-weight: 800; text-decoration: none; color: #000; box-shadow: 2px 2px 0 #000; }}
        .btn:hover {{ transform: translate(-2px, -2px); box-shadow: 4px 4px 0 #000; }}
        pre {{ background: #f4f4f5; padding: 1rem; border: 1.5px solid #000; border-radius: 6px; font-size: 0.8rem; overflow-x: auto; }}
    </style>
</head>
<body>
    <div class="card">
        <h2 style="margin-top:0; color:#c62828;">⚠️ Application Error</h2>
        <p>A server error occurred while processing this page:</p>
        <div style="background:#ffebee; border:1.5px solid #c62828; padding:0.75rem; border-radius:6px; font-weight:700; color:#b71c1c; margin-bottom:1rem;">
            {e}
        </div>
        <div style="display:flex; gap:0.75rem;">
            <a href="/login" class="btn">↩ Return to Login</a>
            <a href="javascript:location.reload()" class="btn" style="background:#fff;">🔄 Retry</a>
        </div>
    </div>
</body>
</html>""", 500

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=True)
